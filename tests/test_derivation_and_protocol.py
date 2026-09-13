from pathlib import Path
import hashlib
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile

from content_discovery_capture.derivation import derive
from content_discovery_capture.verification import verify
from content_discovery_capture.sources.filesystem import FilesystemAdapter
from content_discovery_capture.domain import Scope, Budget, CaptureError
from content_discovery_capture.application import Application
from content_discovery_capture.interfaces.review import decide


class DerivationTests(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec("pdfplumber") and importlib.util.find_spec("reportlab"), "PDF parser and fixture generator required")
    def test_pdf_page_evidence(self):
        from reportlab.pdfgen.canvas import Canvas
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            path = root / "pages.pdf"
            canvas = Canvas(str(path))
            canvas.drawString(70, 700, "First page evidence")
            canvas.showPage()
            canvas.drawString(70, 700, "Second page evidence")
            canvas.save()
            text = self.run_copy(root, path, "application/pdf")
            self.assertIn("## Page 2", text)
            self.assertIn("Second page evidence", text)

    def run_copy(self, root, path, media):
        output = derive(path, media, "text", root / (path.stem + "-derived"), 30, 1000000)
        self.assertTrue(output["limitations"])
        file = output["files"][-1]
        data = (root / (path.stem + "-derived") / file["name"]).read_bytes()
        self.assertEqual(hashlib.sha256(data).hexdigest(), file["sha256"])
        return data.decode()

    @unittest.skipUnless(importlib.util.find_spec("docx"), "Optional DOCX converter not installed")
    def test_docx_table_order(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            document = Document()
            document.add_paragraph("Before table")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
            document.add_paragraph("After table")
            path = root / "ordered.docx"
            document.save(path)
            text = self.run_copy(root, path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            self.assertLess(text.index("Before table"), text.index("Table evidence"))
            self.assertLess(text.index("Table evidence"), text.index("After table"))

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "Optional PPTX converter not installed")
    def test_pptx_text_and_speaker_notes(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            slides = Presentation()
            slide = slides.slides.add_slide(slides.slide_layouts[1])
            slide.shapes.title.text = "Lesson title"
            slide.notes_slide.notes_text_frame.text = "Presenter evidence"
            path = root / "lesson.pptx"
            slides.save(path)
            text = self.run_copy(root, path, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
            self.assertIn("Lesson title", text)
            self.assertIn("Presenter evidence", text)

    def test_odt_and_html_reading_copy(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            path = root / "lesson.odt"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("content.xml", '<office:document xmlns:office="urn:office" xmlns:text="urn:text"><text:p>ODT evidence</text:p></office:document>')
            self.assertIn("ODT evidence", self.run_copy(root, path, "application/vnd.oasis.opendocument.text"))
            html = root / "lesson.html"
            html.write_text('<h1>HTML title</h1><script>Do not include</script><p><a href="/reading">Reading</a></p>')
            text = self.run_copy(root, html, "text/html")
            self.assertIn("HTML title", text)
            self.assertNotIn("Do not include", text)

    @unittest.skipUnless(os.environ.get("CDC_TEST_AUDIO") and os.environ.get("CDC_TRANSCRIPTION_MODEL"), "Explicit local test audio and model required")
    def test_real_local_transcript_hash_timestamps_and_text(self):
        path = Path(os.environ["CDC_TEST_AUDIO"]).resolve()
        with tempfile.TemporaryDirectory() as temp:
            output = derive(path, "audio/flac", "transcript", Path(temp) / "derived", 180, 1000000)
            self.assertEqual({"transcript.md", "transcript.srt", "transcript.json"}, {f["name"] for f in output["files"]})
            result = json.loads((Path(temp) / "derived/transcript.json").read_text())
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), result["source_sha256"])
            self.assertGreater(len(result["segments"]), 0)
            for segment in result["segments"]:
                self.assertGreater(segment["end"], segment["start"])
                self.assertLessEqual(segment["end"], result["duration"] + 1)
            expected = os.environ.get("CDC_TEST_EXPECT_TEXT")
            if expected:
                text = " ".join(s["text"] for s in result["segments"]).casefold()
                self.assertIn(expected.casefold(), text)

    @unittest.skipUnless(os.environ.get("CDC_TEST_AUDIO") and os.environ.get("CDC_TRANSCRIPTION_MODEL"), "Explicit local test audio and model required")
    def test_four_selected_recordings_keep_separate_transcript_lineage(self):
        # Four placements of a public/local test recording exercise lineage, not a course replay.
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            folder = root / "recordings"
            folder.mkdir()
            audio = Path(os.environ["CDC_TEST_AUDIO"]).read_bytes()
            for i in range(4):
                (folder / f"recording-{i}.flac").write_bytes(audio)
            app = Application(root / "project")
            try:
                app.register_source("Test recordings", "filesystem", str(folder))
                review = app.prepare_refresh()
                decide(app.store, review["id"], True)
                snapshot = app.refresh(review["id"])["sources"][0]["snapshot_id"]
                capture = app.prepare_capture([snapshot], [r["selection_id"] for r in app.index([snapshot])["resources"]], derivatives=["transcript"], max_seconds=240)
                decide(app.store, capture["id"], True)
                self.assertEqual("COMPLETE", app.capture(capture["id"], steps=20)["state"])
                originals = [a for a in app.store.all("artefact") if a["role"] == "original"]
                transcripts = [a for a in app.store.all("artefact") if a.get("derivation") == "transcript"]
                self.assertEqual(4, len(originals))
                self.assertEqual(12, len(transcripts))
                self.assertEqual({a["version_id"] for a in originals}, {a["inputs"][0]["version_id"] for a in transcripts})
                self.assertEqual(1, len({a["sha256"] for a in originals}))
            finally:
                app.close()


class ProtocolAndTraversalTests(unittest.TestCase):
    def test_scope_rejects_truthy_strings_and_registry_exposes_boundaries(self):
        with self.assertRaises(CaptureError):
            Scope.from_dict({"roots": ["https://example.org/course"], "allow_private_network": "false"})
        with self.assertRaises(CaptureError):
            Scope.from_dict({"roots": "https://example.org/course"})
        with tempfile.TemporaryDirectory() as temp:
            app = Application(temp)
            try:
                app.register_source("Course", "web", "https://example.org/course")
                registry = app.registry()["sources"][0]
                self.assertEqual(["https://example.org/course"], registry["scope"]["roots"])
                self.assertEqual([], registry["discovery_history"])
                self.assertEqual(registry["scope"], app.prepare_refresh()["payload"]["sources"][0]["scope"])
                with self.assertRaises(CaptureError):
                    app.configure_host_browser({"inspect": "false"})
            finally:
                app.close()

    def test_stdio_roundtrip_and_no_approval_tool(self):
        messages = [{"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {"protocolVersion": "2025-11-25"}},
                    {"jsonrpc": "2.0", "method": "notifications/initialized"},
                    {"jsonrpc": "2.0", "id": 2, "method": "tools/list"},
                    {"jsonrpc": "2.0", "id": 3, "method": "tools/call", "params": {"name": "registry", "arguments": {}}},
                    {"jsonrpc": "2.0", "id": 4, "method": "tools/call", "params": {"name": "approve", "arguments": {}}}]
        with tempfile.TemporaryDirectory() as temp:
            run = subprocess.run([sys.executable, "-m", "content_discovery_capture.interfaces.cli", "--project", temp, "serve"],
                                 input="\n".join(json.dumps(m) for m in messages) + "\n", capture_output=True, text=True, timeout=20)
            self.assertEqual(0, run.returncode, run.stderr)
            responses = [json.loads(line) for line in run.stdout.splitlines()]
            self.assertEqual([1, 2, 3, 4], [r["id"] for r in responses])
            self.assertFalse(any("approv" in t["name"] for t in responses[1]["result"]["tools"]))
            self.assertFalse(responses[2]["result"]["isError"])
            self.assertTrue(responses[3]["result"]["isError"])

    def test_large_directory_lexical_checkpoint_has_no_omissions(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            folder = root / "input"
            folder.mkdir()
            for i in reversed(range(600)):
                (folder / f"item-{i:03}.txt").touch()
            adapter = FilesystemAdapter(root / "project")
            scope, cursor, seen = Scope((str(folder),)), {"location": str(folder), "depth": 0}, []
            while cursor:
                batch = adapter.discover({}, scope, cursor, Budget())
                self.assertLessEqual(len(batch.frontier), 257)
                seen.extend(c["location"] for c in batch.frontier if c["location"] != str(folder))
                cursor = next((c for c in batch.frontier if c["location"] == str(folder)), None)
            self.assertEqual(600, len(seen))
            self.assertEqual(600, len(set(seen)))
            app = Application(root / "project")
            try:
                app.register_source("Large folder", "filesystem", str(folder))
                review = app.prepare_refresh()
                decide(app.store, review["id"], True)
                result = app.refresh(review["id"], steps=1000)
                self.assertEqual("COMPLETE", result["sources"][0]["state"])
                self.assertEqual(600, result["sources"][0]["resources"])
            finally:
                app.close()


if __name__ == "__main__":
    unittest.main()
