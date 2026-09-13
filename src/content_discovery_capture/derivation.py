"""Isolated, bounded transformation worker. Never downloads transcription models."""
from __future__ import annotations
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin
import importlib.util
import json
import os
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET

from .domain import CaptureError, safe_text


class MarkdownParser(HTMLParser):
    def __init__(self, base):
        super().__init__(convert_charrefs=True)
        self.base, self.parts, self.skip, self.hrefs = base, [], 0, []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in ("script", "style", "template"):
            self.skip += 1
        if self.skip:
            return
        if tag in ("p", "div", "section", "article", "tr", "ul", "ol"):
            self.parts.append("\n\n")
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.parts.append("\n\n" + "#" * int(tag[1]) + " ")
        if tag == "li":
            self.parts.append("\n- ")
        if tag in ("td", "th"):
            self.parts.append(" | ")
        if tag == "br":
            self.parts.append("\n")
        if tag == "a":
            self.hrefs.append(safe_text(urljoin(self.base, attrs.get("href", ""))))
            self.parts.append("[")
        if tag == "img":
            self.parts.append(f"![{attrs.get('alt', '')}]({safe_text(urljoin(self.base, attrs.get('src', '')))})")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "template"):
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag == "a" and self.hrefs:
            self.parts.append("](" + self.hrefs.pop() + ")")
        if tag in ("p", "div", "section", "article", "tr") or tag.startswith("h") and len(tag) == 2:
            self.parts.append("\n\n")

    def handle_data(self, text):
        if not self.skip:
            self.parts.append(text)


def supported(media, operation):
    if operation == "transcript":
        model = os.environ.get("CDC_TRANSCRIPTION_MODEL", "")
        return media.startswith(("audio/", "video/")) and bool(model) and Path(model).is_dir() and importlib.util.find_spec("faster_whisper") is not None
    if "html" in media or media.startswith("text/") or "opendocument" in media:
        return True
    return any(marker in media and importlib.util.find_spec(module) is not None for marker, module in
               [("pdf", "pdfplumber"), ("wordprocessingml", "docx"), ("presentationml", "pptx")])


def derive(input_path, media, operation, destination, timeout, max_bytes, base=""):
    if not supported(media, operation):
        raise CaptureError("This derivation is unavailable for the representation or installed capabilities.")
    destination.mkdir(exist_ok=True)
    request = destination / "request.json"
    request.write_text(json.dumps({"input": str(input_path), "media": media, "operation": operation,
                                   "destination": str(destination), "base": base, "max_output_bytes": max_bytes}))
    try:
        result = subprocess.run([sys.executable, "-m", "content_discovery_capture.derivation", str(request)],
                                timeout=timeout, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except subprocess.TimeoutExpired:
        raise CaptureError("Derivation reached the approved time budget.") from None
    if result.returncode:
        raise CaptureError("Derivation failed; the original remains preserved.")
    output = json.loads((destination / "result.json").read_text())
    if sum((destination / f["name"]).stat().st_size for f in output["files"]) > max_bytes:
        raise CaptureError("Derived output exceeds the approved byte budget.")
    return output


def worker(request):
    src, out = Path(request["input"]), Path(request["destination"])
    try:
        import resource
        limit = request.get("max_output_bytes", 1024**3)
        resource.setrlimit(resource.RLIMIT_FSIZE, (limit, limit))
    except ImportError:
        pass
    media, operation = request["media"], request["operation"]
    limitations, files = [], []
    if operation == "transcript":
        from faster_whisper import WhisperModel
        model = WhisperModel(os.environ["CDC_TRANSCRIPTION_MODEL"], device="cpu", compute_type="int8", local_files_only=True)
        segments, info = model.transcribe(str(src), word_timestamps=True, vad_filter=True)
        segments = [{"start": s.start, "end": s.end, "text": s.text, "avg_logprob": s.avg_logprob,
                     "no_speech_prob": s.no_speech_prob, "words": [w._asdict() for w in (s.words or [])]} for s in segments]
        text = "\n\n".join(f"[{s['start']:.2f}–{s['end']:.2f}] {s['text'].strip()}" for s in segments)
        flagged = [i for i, s in enumerate(segments) if s["avg_logprob"] < -1 or s["no_speech_prob"] > .6 or s["end"] <= s["start"]]
        import hashlib
        from collections import Counter
        repetitions = Counter(s["text"].strip().casefold() for s in segments)
        flagged = sorted(set(flagged + [i for i, s in enumerate(segments) if repetitions[s["text"].strip().casefold()] >= 3]))
        with src.open("rb") as stream:
            source_hash = hashlib.file_digest(stream, "sha256").hexdigest()
        (out / "transcript.json").write_text(json.dumps({"segments": segments, "duration": info.duration,
            "source_sha256": source_hash,
            "language": info.language, "flagged_segments": flagged, "model": Path(os.environ["CDC_TRANSCRIPTION_MODEL"]).name,
            "configuration": {"device": "cpu", "compute_type": "int8", "word_timestamps": True, "vad_filter": True}}, indent=2))
        def stamp(seconds):
            ms = round(seconds * 1000)
            return f"{ms // 3600000:02}:{ms // 60000 % 60:02}:{ms // 1000 % 60:02},{ms % 1000:03}"
        (out / "transcript.srt").write_text("\n\n".join(f"{i+1}\n{stamp(s['start'])} --> {stamp(s['end'])}\n{s['text'].strip()}" for i, s in enumerate(segments)))
        files += [{"name": "transcript.json", "media_type": "application/json"}, {"name": "transcript.srt", "media_type": "text/plain"}]
        limitations.append("Machine transcript; not certified or fully reviewed. Consult the recording for exact quotations.")
    elif "html" in media:
        parser = MarkdownParser(request["base"])
        parser.feed(src.read_text(errors="replace"))
        text = "".join(parser.parts)
        limitations.append("Rendered layout, interactions and unselected assets may not be represented")
    elif "wordprocessingml" in media:
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        document = Document(src)
        parts = []
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                parts.append(item.text)
            elif isinstance(item, Table):
                parts.extend(" | ".join(c.text for c in row.cells) for row in item.rows)
        text = "\n\n".join(parts)
        limitations.append("Original is required for layout, fields, images and embedded objects")
    elif "presentationml" in media:
        from pptx import Presentation
        parts = []
        for number, slide in enumerate(Presentation(src).slides, 1):
            parts.append(f"## Slide {number}")
            for shape in sorted(slide.shapes, key=lambda s: (s.top, s.left)):
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    parts.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
        text = "\n\n".join(parts)
        limitations.append("Original is required for diagrams, images, animation and layout")
    elif "pdf" in media:
        import pdfplumber
        with pdfplumber.open(src) as pdf:
            text = "\n\n".join(f"## Page {i}\n\n{page.extract_text() or '[No extractable text]'}" for i, page in enumerate(pdf.pages, 1))
        limitations.append("Text extraction may omit scanned text, figures and layout; no OCR was performed")
    elif "opendocument" in media:
        with zipfile.ZipFile(src) as archive:
            xml = ET.fromstring(archive.read("content.xml"))
        text = "\n\n".join("".join(node.itertext()) for node in xml.iter() if node.tag.endswith(("}p", "}h")))
        limitations.append("Original retains layout and embedded objects")
    else:
        text = src.read_text(errors="replace")
    if not text.strip():
        raise ValueError("Empty derived text")
    name = "transcript.md" if operation == "transcript" else "reading.md"
    (out / name).write_text(safe_text(text.strip()) + "\n")
    files.append({"name": name, "media_type": "text/markdown"})
    import hashlib
    for file in files:
        data = (out / file["name"]).read_bytes()
        file["sha256"], file["size"] = hashlib.sha256(data).hexdigest(), len(data)
    (out / "result.json").write_text(json.dumps({"files": files, "limitations": limitations, "operation": operation, "worker_version": "1"}))


if __name__ == "__main__":
    worker(json.loads(Path(sys.argv[1]).read_text()))
